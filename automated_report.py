import os
from datetime import datetime

import pandas as pd
import matplotlib.pyplot as plt

from colorama import Fore, Style, init
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as ExcelImage


# Enable terminal colours on Windows
init(autoreset=True)

OUTPUT_FILE = "Automated_Report.xlsx"
CHART_FILE = "sales_chart.png"

# Terminal text colours
RESET = Style.RESET_ALL
BLUE = Fore.BLUE
GREEN = Fore.GREEN
YELLOW = Fore.YELLOW
PURPLE = Fore.MAGENTA
CYAN = Fore.CYAN
RED = Fore.RED
BOLD = Style.BRIGHT

# Excel row colours
EXCEL_COLOURS = [
    "D9EAF7", "E2F0D9", "FFF2CC", "E4DFEC",
    "DDEBF7", "FCE4D6", "F4CCCC", "D9EAD3"
]

# Bar graph colours
BAR_COLOURS = [
    "red", "blue", "green", "orange", "purple",
    "cyan", "pink", "gold", "brown", "lime"
]


# 1. Enter product details
def get_sales_data():
    products = []

    print(BOLD + BLUE + "\n" + "=" * 55)
    print("       AUTOMATED EXCEL REPORT GENERATOR")
    print("=" * 55 + RESET)

    while True:
        print(CYAN + "\nEnter Product Details")

        while True:
            name = input("Enter product name: ").strip()
            if name:
                break
            print(RED + "Product name cannot be empty.")

        while True:
            try:
                quantity = int(input("Enter quantity sold: "))
                if quantity > 0:
                    break
                print(RED + "Quantity must be greater than zero.")
            except ValueError:
                print(RED + "Enter a valid whole number.")

        while True:
            try:
                price = float(input("Enter selling price per item: "))
                if price >= 0:
                    break
                print(RED + "Price cannot be negative.")
            except ValueError:
                print(RED + "Enter a valid number.")

        while True:
            try:
                cost_price = float(input("Enter cost price per item: "))
                if cost_price >= 0:
                    break
                print(RED + "Cost price cannot be negative.")
            except ValueError:
                print(RED + "Enter a valid number.")

        total_sales = quantity * price
        total_cost = quantity * cost_price
        profit = total_sales - total_cost

        products.append({
            "Product": name,
            "Quantity": quantity,
            "Selling Price": price,
            "Cost Price": cost_price,
            "Total Sales": total_sales,
            "Total Cost": total_cost,
            "Profit": profit,
            "Date": datetime.now().strftime("%Y-%m-%d")
        })

        colour = [
            BLUE, GREEN, YELLOW, PURPLE, CYAN, RED
        ][(len(products) - 1) % 6]

        print(colour + "\nProduct added successfully!")
        print(f"Product     : {name}")
        print(f"Total Sales : Rs. {total_sales:,.2f}")
        print(f"Total Cost  : Rs. {total_cost:,.2f}")
        print(f"Profit      : Rs. {profit:,.2f}" + RESET)

        choice = input(
            "\nDo you want to add another product? (yes/no): "
        ).strip().lower()

        if choice not in ("yes", "y"):
            break

    return products


# 2. Calculate sales summary
def calculate_summary(df):
    total_sales = df["Total Sales"].sum()
    total_cost = df["Total Cost"].sum()
    total_profit = df["Profit"].sum()

    profit_margin = (
        total_profit / total_sales * 100
        if total_sales > 0 else 0
    )

    return {
        "Total Products": len(df),
        "Total Quantity Sold": int(df["Quantity"].sum()),
        "Total Sales": total_sales,
        "Total Cost": total_cost,
        "Total Profit": total_profit,
        "Profit Margin (%)": profit_margin
    }


# 3. Display all products in different terminal colours
def display_all_items(df):
    print(BOLD + BLUE + "\n" + "=" * 60)
    print("               ALL PRODUCT DETAILS")
    print("=" * 60 + RESET)

    colours = [BLUE, GREEN, YELLOW, PURPLE, CYAN, RED]

    for index, row in df.iterrows():
        colour = colours[index % len(colours)]

        print(colour + "-" * 50)
        print(f"Product       : {row['Product']}")
        print(f"Quantity      : {row['Quantity']}")
        print(f"Selling Price : Rs. {row['Selling Price']:,.2f}")
        print(f"Cost Price    : Rs. {row['Cost Price']:,.2f}")
        print(f"Total Sales   : Rs. {row['Total Sales']:,.2f}")
        print(f"Total Cost    : Rs. {row['Total Cost']:,.2f}")
        print(f"Profit        : Rs. {row['Profit']:,.2f}")
        print(f"Date          : {row['Date']}")
        print("-" * 50 + RESET)


# 4. Display summary
def display_summary(summary):
    print(BOLD + PURPLE + "\n" + "=" * 55)
    print("                  SALES SUMMARY")
    print("=" * 55 + RESET)

    for metric, value in summary.items():
        if metric in ("Total Sales", "Total Cost", "Total Profit"):
            print(GREEN + f"{metric:<25}: Rs. {value:,.2f}")

        elif metric == "Profit Margin (%)":
            print(YELLOW + f"{metric:<25}: {value:.2f}%")

        else:
            print(CYAN + f"{metric:<25}: {value}")

    print(PURPLE + "=" * 55 + RESET)


# 5. Create and display a multicolour bar graph
def create_sales_chart(df):
    chart_data = df.groupby("Product")["Total Sales"].sum()

    colours = [
        BAR_COLOURS[i % len(BAR_COLOURS)]
        for i in range(len(chart_data))
    ]

    fig, ax = plt.subplots(figsize=(12, 7))

    bars = ax.bar(
        chart_data.index,
        chart_data.values,
        color=colours,
        edgecolor="black",
        linewidth=0.8
    )

    # Display the sales amount above every bar
    for bar in bars:
        height = bar.get_height()

        ax.annotate(
            f"Rs. {height:,.0f}",
            xy=(bar.get_x() + bar.get_width() / 2, height),
            xytext=(0, 5),
            textcoords="offset points",
            ha="center",
            va="bottom",
            fontsize=9,
            fontweight="bold"
        )

    ax.set_title(
        "Product-wise Sales Report",
        fontsize=16,
        fontweight="bold"
    )
    ax.set_xlabel("Product Name", fontsize=12)
    ax.set_ylabel("Total Sales (Rs.)", fontsize=12)

    ax.tick_params(axis="x", labelrotation=30)
    ax.grid(axis="y", linestyle="--", alpha=0.3)
    ax.set_axisbelow(True)

    plt.setp(ax.get_xticklabels(), ha="right")
    fig.tight_layout()

    # Save the colourful graph
    fig.savefig(CHART_FILE, dpi=150, bbox_inches="tight")

    # Display the graph on screen
    plt.show()
    plt.close(fig)

    print(GREEN + "\nMulticolour bar graph displayed.")
    print(f"Chart saved: {os.path.abspath(CHART_FILE)}")


# 6. Format Excel worksheets with different row colours
def format_worksheet(ws):
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
        size=11
    )

    # Format header
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Apply a different background colour to each row
    for row_number in range(2, ws.max_row + 1):
        colour = EXCEL_COLOURS[
            (row_number - 2) % len(EXCEL_COLOURS)
        ]

        for cell in ws[row_number]:
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=colour
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

    # Set column widths
    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        ws.column_dimensions[column_letter].width = min(
            max_length + 4, 30
        )

    # Format money columns
    headers = [cell.value for cell in ws[1]]

    for name in [
        "Selling Price", "Cost Price",
        "Total Sales", "Total Cost", "Profit"
    ]:
        if name in headers:
            column = headers.index(name) + 1

            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row, column)
                cell.number_format = '"Rs. "#,##0.00'

                if name == "Profit" and isinstance(
                    cell.value, (int, float)
                ):
                    colour = (
                        "008000" if cell.value >= 0
                        else "FF0000"
                    )
                    cell.font = Font(
                        bold=True,
                        color=colour
                    )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 25


# 7. Generate formatted Excel report
def create_excel_report(df, summary):
    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl"
    ) as writer:

        df.to_excel(
            writer,
            sheet_name="Sales Data",
            index=False
        )

        summary_df = pd.DataFrame(
            list(summary.items()),
            columns=["Metric", "Value"]
        )

        summary_df.to_excel(
            writer,
            sheet_name="Summary",
            index=False,
            startrow=1
        )

    workbook = load_workbook(OUTPUT_FILE)

    sales_ws = workbook["Sales Data"]
    summary_ws = workbook["Summary"]

    format_worksheet(sales_ws)
    format_worksheet(summary_ws)

    # Add a colourful report title
    summary_ws["A1"] = "AUTOMATED SALES REPORT"
    summary_ws["A1"].font = Font(
        bold=True,
        size=16,
        color="FFFFFF"
    )
    summary_ws["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="7030A0"
    )

    summary_ws.merge_cells("A1:B1")
    summary_ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )
    summary_ws.row_dimensions[1].height = 30

    summary_ws.column_dimensions["A"].width = 28
    summary_ws.column_dimensions["B"].width = 22

    # Format summary values
    for row in range(3, summary_ws.max_row + 1):
        metric = summary_ws.cell(row, 1).value
        value_cell = summary_ws.cell(row, 2)

        if metric in ("Total Sales", "Total Cost", "Total Profit"):
            value_cell.number_format = '"Rs. "#,##0.00'

        elif metric == "Profit Margin (%)":
            value_cell.number_format = '0.00"%"'

    # Embed the multicolour bar graph in the summary sheet
    if os.path.exists(CHART_FILE):
        chart = ExcelImage(CHART_FILE)
        chart.width = 650
        chart.height = 350
        summary_ws.add_image(chart, "D2")

    workbook.save(OUTPUT_FILE)

    print(GREEN + "\nExcel report generated successfully!")
    print(f"Excel file: {os.path.abspath(OUTPUT_FILE)}")


# 8. Main program
def main():
    try:
        products = get_sales_data()

        if not products:
            print(RED + "No product details entered.")
            return

        df = pd.DataFrame(products)
        summary = calculate_summary(df)

        # Display product details and summary on screen
        display_all_items(df)
        display_summary(summary)

        # Display colourful chart and create Excel report
        create_sales_chart(df)
        create_excel_report(df, summary)

        print(BOLD + GREEN + "\n" + "=" * 55)
        print("       PROJECT COMPLETED SUCCESSFULLY!")
        print("=" * 55 + RESET)

        print("1. Product details displayed.")
        print("2. Terminal colours applied.")
        print("3. Multicolour bar graph generated.")
        print("4. Excel rows formatted with different colours.")
        print("5. Sales summary and graph added to Excel.")

    except PermissionError:
        print(
            RED + "\nPermission error: Close the Excel report "
            "if it is open, then run the program again."
        )

    except Exception as error:
        print(RED + f"\nAn error occurred: {error}")


if __name__ == "__main__":
    main()
